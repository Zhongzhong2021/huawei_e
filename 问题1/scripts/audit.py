"""Rebuild all inputs and native clocks directly from contest files."""
import argparse,json
from concurrent.futures import ThreadPoolExecutor
from fractions import Fraction
from pathlib import Path
import numpy as np
from openpyxl import load_workbook
from common import digest,run,write

def inspect(path):
    raw=run(['ffprobe','-v','error','-show_streams','-show_format','-show_frames','-show_entries',
        'stream=index,codec_type,sample_rate,channels,width,height,time_base:format=duration:frame=media_type,best_effort_timestamp,pkt_duration,nb_samples',
        '-of','json',str(path)])
    p=json.loads(raw);clocks={}
    for kind in ['audio','video']:
        streams=[s for s in p['streams'] if s['codec_type']==kind];assert len(streams)==1
        s=streams[0];frames=[f for f in p['frames'] if f['media_type']==kind];tb=Fraction(s['time_base'])
        ticks=[int(f['best_effort_timestamp']) for f in frames]
        assert ticks and all(b>a for a,b in zip(ticks,ticks[1:]))
        lengths=[int(f['nb_samples']) if kind=='audio' else int(f['pkt_duration']) for f in frames]
        dt=Fraction(1,int(s['sample_rate'])) if kind=='audio' else tb
        starts=[t*tb for t in ticks];ends=[t+l*dt for t,l in zip(starts,lengths)]
        assert all(e>a for a,e in zip(starts,ends)) and all(b>=a for a,b in zip(ends,starts[1:]))
        clocks[kind]={**s,'pts_ticks':ticks,'lengths':lengths,'length_unit':'samples' if kind=='audio' else 'ticks',
            'start_s':float(starts[0]),'end_s':float(ends[-1]),'positive_gaps':sum(b>a for a,b in zip(ends,starts[1:]))}
    a=clocks['audio'];assert a['positive_gaps']==0,'Explicit handling required for audio gaps'
    pcm=np.frombuffer(run(['ffmpeg','-nostdin','-v','error','-threads','1','-i',str(path),'-map','0:a:0','-c:a','pcm_f32le','-f','f32le','-']),dtype='<f4')
    assert len(pcm)%a['channels']==0 and len(pcm)//a['channels']==sum(a['lengths']) and np.isfinite(pcm).all()
    return {'sha256':digest(path),'container_duration_s':float(p['format']['duration']),'audio':a,'video':clocks['video'],
        'pcm_samples':len(pcm)//a['channels'],'zero_audio':not bool(np.any(pcm)),'decode':'FFmpeg default edit-list handling'}

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--source',type=Path,required=True);args=ap.parse_args()
    books=list(args.source.rglob('label-100.xlsx'));assert len(books)==1
    wb=load_workbook(books[0],read_only=True,data_only=True);sheet=wb[wb.sheetnames[0]]
    it=sheet.iter_rows(values_only=True);header=next(it);labels=[]
    for row in it:
        if not any(v is not None for v in row):continue
        r=dict(zip(header,row));r={str(k).strip():v for k,v in r.items() if k is not None}
        r['video_id']=str(r['video_id']);r['clip_id']=str(int(r['clip_id']));r['sample_id']=r['video_id']+'$_$'+r['clip_id']
        assert isinstance(r['text'],str) and -3<=float(r['label'])<=3
        expected='Negative' if float(r['label'])<0 else 'Positive' if float(r['label'])>0 else 'Neutral';assert r['annotation']==expected
        labels.append(r)
    assert len(labels)==100 and len({r['sample_id'] for r in labels})==100
    videos=list(args.source.rglob('*.mp4'));paths={p.parent.name+'$_$'+p.stem:p for p in videos}
    assert len(videos)==len(paths)==100 and set(paths)=={r['sample_id'] for r in labels}
    def one(r):
        path=paths[r['sample_id']];before=digest(path);media=inspect(path);assert before==digest(path)==media['sha256']
        return {**r,**media,'path':str(path.relative_to(args.source))}
    with ThreadPoolExecutor(max_workers=4) as pool:items=list(pool.map(one,sorted(labels,key=lambda r:r['sample_id'])))
    write('data/inputs.json',{'source_root':str(args.source.resolve()),'label_sha256':digest(books[0]),'samples':items})
    write('data/audit.json',{'samples':len(items),'videos':len({r['video_id'] for r in items}),'zero_audio':[r['sample_id'] for r in items if r['zero_audio']],
        'observed_spans_s':sum(max(r['audio']['end_s'],r['video']['end_s'])-min(r['audio']['start_s'],r['video']['start_s']) for r in items),
        'raw_decoding_and_mappings_checked':True,'semantic_correctness':'not_established'})
    print('Rebuilt 100 inputs, native clocks and original text; labels retained without alteration.')
if __name__=='__main__':main()
