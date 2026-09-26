"""Regenerate the formal paper, figures and all twenty explanation cards from frozen results."""
from pathlib import Path
import csv, json
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
PAPER = ROOT/'paper'
FIG = PAPER/'figures'
FIG.mkdir(exist_ok=True)

def read(p): return json.loads((ROOT/p).read_text())
def rows(p):
    with (ROOT/p).open(encoding='utf-8-sig',newline='') as f: return list(csv.DictReader(f))
def records(p): return [json.loads(x) for x in (ROOT/p).read_text().splitlines() if x]
def save(fig,name):
    fig.savefig(FIG/(name+'.png'),dpi=200,bbox_inches='tight')
    fig.savefig(FIG/(name+'.svg'),bbox_inches='tight')
    plt.close(fig)

MET = read('results/uncertainty/metrics.json')
CON = read('results/contributions/contribution_summary.json')['summary']
CMP = rows('results/model_comparison.csv')
MEDIA = {x['id']:x for x in records('results/media_mapping.jsonl')}
EXP = {x['id']:x for x in records('results/uncertainty/special_explanations.jsonl')}
PRED = {x['id']:x for x in rows('results/uncertainty/special_predictions.csv')}
LABEL = {x['sample_id']:x for x in rows('results/public_label_audit/attachment4_public_labels.csv')}
SHARE = {x['id']:x for x in rows('results/contributions/attachment4_contributions.csv')}
def strength(v):
    v=float(v)
    return f'{v:+.1e}' if 0<abs(v)<1e-4 else f'{v:+.4f}'

CN = ['负向','中性','正向']
COLORS = ['#277DA8','#D89635','#5B9982']
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10,'axes.spines.top':False,'axes.spines.right':False,'svg.fonttype':'none'})

# Network: boxes label operations, arrows follow the actual dependency graph.
fig, ax=plt.subplots(figsize=(12,7.2));ax.set_xlim(0,12);ax.set_ylim(-.18,8);ax.axis('off')
def box(x,y,w,h,t,c='#EAF2F7'):
    ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=0.09',fc=c,ec='#597080',lw=1))
    ax.text(x+w/2,y+h/2,t,ha='center',va='center',fontsize=10)
def arrow(a,b):ax.annotate('',xy=b,xytext=a,arrowprops=dict(arrowstyle='->',color='#597080',lw=1.2))
for i,(name,enc) in enumerate([('T: token IDs + masks','BERT (12 layers)\ncontextual rows: 768'),('A: aligned features','Standardized acoustic\nfeatures: 74'),('V: aligned features','Standardized visual\nfeatures: 35')]):
    x=.1+4*i;box(x,7,3.7,.65,name);box(x,5.75,3.7,.85,enc);arrow((x+1.85,7),(x+1.85,6.6));box(x,4.45,3.7,.85,'Local encoder: D_m -> 64 -> 64\nE(x) - E(0), window = 1');arrow((x+1.85,5.75),(x+1.85,5.3))
box(.15,2.8,4.0,1,'Three single branches\n64 -> 4; masked mean pooling')
box(4.55,2.8,3.45,1,'TA / TV / AV pairs\nrank-8 bilinear -> 4')
box(8.4,2.8,3.45,1,'Unimodal scores + activity\n8 -> 1 log variance / modality')
for x in [1.95,5.95,9.95]:arrow((x,4.45),(2.15,3.8));arrow((x,4.45),(6.3,3.8))
ax.plot([4.15,4.32,4.32,10.1],[3.3,3.3,4.03,4.03],color='#597080',lw=1.2,ls='--');arrow((10.1,4.03),(10.1,3.8));box(8.4,1.35,3.45,1,'Normalized precision w\nsingle scale K*w\npair scale sqrt(s_m*s_n)');arrow((10.1,2.8),(10.1,2.35))
box(.15,1.35,7.6,1,'Weighted local single + pair terms + bias\nq (regression) and 3 classification logits');arrow((2.15,2.8),(2.15,2.35));arrow((6.3,2.8),(6.3,2.35));arrow((8.4,1.85),(7.75,1.85))
box(.15,0,5.5,.85,'Prediction: neutral logit - 0.4; argmax\n3*tanh(q), then class-consistent intensity','#E8F2ED')
box(6.05,0,5.8,.85,'Native explanation: calibrated logit margin\nlocal evidence + equal allocation of pair terms','#FFF2DE')
arrow((3.1,1.35),(2.9,.85));arrow((5.2,1.35),(8.95,.85));save(fig,'01_network')

hist=rows('results/uncertainty/history.csv');ep=[int(x['epoch']) for x in hist]
fig,axes=plt.subplots(1,2,figsize=(10,3.5))
axes[0].plot(ep,[float(x['train_loss']) for x in hist],'o-',color=COLORS[0]);axes[0].set(xlabel='Epoch',ylabel='Joint training loss')
axes[1].plot(ep,[float(x['valid_raw_mae']) for x in hist],'o-',color=COLORS[1]);axes[1].set(xlabel='Epoch',ylabel='Validation raw MAE')
for ax in axes:ax.axvline(5,ls='--',color='#555',label='Selected epoch 5');ax.legend();ax.grid(alpha=.2)
fig.tight_layout();save(fig,'02_training')

fig,axes=plt.subplots(1,2,figsize=(9,3.8))
for ax,split in zip(axes,['valid','test']):
    a=np.asarray(MET[split]['final']['confusion_matrix']);ax.imshow(a,cmap='Blues',vmin=0,vmax=300)
    for (i,j),v in np.ndenumerate(a):ax.text(j,i,str(v),ha='center',va='center',color='white' if v>160 else '#172D3D')
    ax.set(xticks=range(3),yticks=range(3),xticklabels=['Negative','Neutral','Positive'],yticklabels=['Negative','Neutral','Positive'],xlabel='Predicted class',ylabel='True class',title=f'{split.capitalize()} (n={a.sum()})')
fig.tight_layout();save(fig,'03_confusion')

valid=rows('results/uncertainty/valid_predictions.csv');y=np.array([float(x['true_intensity']) for x in valid]);yp=np.array([float(x['intensity']) for x in valid])
fig,axes=plt.subplots(1,2,figsize=(9,3.7));axes[0].scatter(y,yp,s=11,alpha=.35,color=COLORS[0]);axes[0].plot([-3,3],[-3,3],'--',color='#777');axes[0].set(xlabel='True intensity',ylabel='Final intensity',xlim=(-3.1,3.1),ylim=(-3.1,3.1));axes[1].hist(yp-y,bins=25,color=COLORS[0],edgecolor='white');axes[1].axvline(0,color='#555',ls='--');axes[1].set(xlabel='Prediction - truth',ylabel='Number of samples');fig.tight_layout();save(fig,'04_regression')

fig,axes=plt.subplots(1,3,figsize=(11,3.8))
for ax,split in zip(axes,['valid','test','special']):
    data=np.array([[CON[split]['classification']['mean_shares'][m] for m in 'TAV'],[CON[split]['regression']['mean_shares'][m] for m in 'TAV'],[CON[split]['mean_precision_weights'][m] for m in 'TAV']])*100
    bottom=np.zeros(3)
    for j,m in enumerate('TAV'):
        ax.bar(range(3),data[:,j],bottom=bottom,color=COLORS[j],label=m)
        for k,v in enumerate(data[:,j]):
            if v>5:ax.text(k,bottom[k]+v/2,f'{v:.1f}',ha='center',va='center',fontsize=8,color='white' if j==0 else '#172D3D')
        bottom+=data[:,j]
    ax.set(xticks=range(3),xticklabels=['Class margin','Regression q','Reliability'],ylim=(0,106),ylabel='Per-sample share, averaged (%)',title=split.capitalize());ax.legend(ncol=3,loc='upper center',bbox_to_anchor=(.5,1.2))
fig.tight_layout();save(fig,'05_contributions')

# Each displayed word is an exact source token mapping; signs retain opposition.
for sid in ['02','14','16']:
    e=EXP[sid];mapping={j:z for z in MEDIA[sid]['entries'] for j in z['source_indices']}
    indices=[(i,s[0]) for i,s in enumerate(e['source_index']) if s[0]>=0]
    vals=[e['local'][i][0] for i,j in indices]
    labels=[f'{j}: {mapping[j]["text"]}' if j in mapping else str(j) for i,j in indices]
    fig,ax=plt.subplots(figsize=(10,max(3,len(vals)*.18)))
    ax.barh(range(len(vals)),vals,color=[COLORS[0] if v>=0 else '#BC6373' for v in vals]);ax.set(yticks=range(len(vals)),yticklabels=labels,xlabel='Signed local contribution to predicted-class logit margin',title=f'Attachment 4 / {sid}: text readout evidence');ax.invert_yaxis();ax.axvline(0,color='#555',lw=.8);ax.tick_params(axis='y',labelsize=8);fig.tight_layout();save(fig,'06_local_'+sid)

# Paper blocks produce both editable Markdown and a fully illustrated Word document.
blocks=[]
def add(kind,value):blocks.append((kind,value))
def para(s):add('p',s)
def head(s):add('h',s)
def table(headers,data):add('table',(headers,data))
def figure(name,caption):add('figure',(name,caption))
def eq(s):add('eq',s)
add('title','复杂场景下多模态情感识别：基于不确定性融合的原生可解释模型')
para('摘要：针对对齐文本、语音与视觉特征的情感极性识别及强度估计，构建由上下文化局部证据、低秩成对交互和可学习可靠性缩放组成的联合模型。文本采用预训练 BERT[1]，音频和视觉采用训练集统计归一化；以模态单项净分数和局部活性预测单模态残差方差，用归一化精度缩放实际参与预测的证据。模型同时返回三分类 logit 与连续强度，通过校准后的类别间隔进行局部和模态级精确核算。附件2验证集、测试集准确率分别为65.11%和68.64%，Macro-F1分别为0.6323和0.6493，最终强度MAE分别为0.5397和0.6006。附件4全20条预测在公开源标签事后核验中准确率80.00%、MAE 0.3390。文本在测试集分类间隔的平均净贡献占84.62%，在回归预激活中的占比为71.14%，表明分类与强度判断利用模态的方式不同。上述解释是网络读出层面的数值核算，不等同于输入删除的因果效应。')
para('关键词：多模态情感分析；不确定性融合；局部证据；成对交互；原生可解释性')
head('1 问题定义、数据与实验约定')
para('问题三要求在三模态输入条件下，既输出情感极性和强度，又量化三模态作用、确定主要参考模态并定位关键证据。本文仅使用附件2对齐版训练，不以完整CMU-MOSEI替代比赛划分。强度标签y∈[−3,3]，y<0、y=0、y>0依次对应负向、中性、正向，编码为0、1、2；不将接近零但非零的标签重定义为中性。')
table(['集合','样本数','负向','中性','正向','用途'],[['训练',3395,967,758,1670,'统计估计与参数学习'],['验证',728,206,184,338,'模型、轮次、决策策略选优'],['附件2测试',727,207,158,362,'冻结预测后的评价'],['附件4',20,7,0,13,'专项解释与公开源事后核验']])
para('附件2特征每条最多50个对齐位置，文本给定特征768维、语音74维、视觉35维，并提供BERT词元编号、注意力掩码及类型编号。主模型重新通过BERT提取上下文化文本表示。剔除特殊词元和填充位置，在内容位置上保留原source_index；语音、视觉全零行作为无有效数值观测处理，缺失值不纳入训练统计。标准化均值和尺度只由训练集估计，验证、测试和附件4复用同一stats.json。')
para('附件4标签由完整公开CMU-MOSEI[4]数据的文本唯一匹配、特征逐值比对以及SDK标签复核获得。20条均为公开源test样本，其中03、07、08、12、15也在附件2测试集中。故不能将两个测试集直接合并视为747条独立测试；附件4没有中性真值，不能据其80%准确率判断中性识别能力。恢复标签发生在已有预测冻结之后；本研究历史上已多次查看附件2测试结果，且最终主模型是在综合性能与题目解释要求后确定，评价属于探索性比较，不宣称全新盲测或独立确认性优势。')
head('2 建模原理与网络结构')
para('借鉴神经可加模型[5]与成对交互建模[6]，采用“局部非线性编码—显式加性读出”的结构。BERT提供语言上下文，局部编码保留证据位置，成对分支表达同位置的跨模态共同作用。可靠性缩放在样本层调整这些读出项，不额外设置无法核算的融合残差。预测及解释共享同一次前向中的实际数值。')
figure('01_network','图1 不确定性融合网络。T/A/V为文本/语音/视觉；每个分支的4维输出对应q及负向、中性、正向logit。')
para('记m∈{T,A,V}，j为有效内容窗口。当前窗口宽度为1；每模态局部编码器由逐行线性映射与GELU、窗口拼接线性映射与GELU构成，投影维度与局部隐层均为64。文本BERT隐状态先作LayerNorm，再按source_index回填局部位置。用相同观测结构下的零输入编码作为参考，得到下式。参考为特征坐标基准，不代表真实中性情绪。')
eq('dₘⱼ = Eₘ(xₘⱼ, oₘⱼ) − Eₘ(0, oₘⱼ)')
para('单模态读出为无偏置64→4线性映射。均值汇聚权重ωₘⱼ正比于该窗口内有效观测数；成对权重ωₘₙⱼ以两端权重几何均值构造，限制在共同有效窗口，再按样本归一化。TA、TV、AV各使用独立的秩8双线性投影，分别输出4维交互。')
eq('aₘⱼ = ωₘⱼ Wₘdₘⱼ；  iₘₙⱼ = ωₘₙⱼ Vₘₙ[(Pₘₙdₘⱼ) ⊙ (Qₘₙdₙⱼ)]')
para('P、Q、V均不含偏置；某一模态没有有效观测时，该单项与关联交互严格为零。BERT使单个文本读出含有全句上下文，因此不能把一个局部值解释为只由该词自身产生的独立效应。成对模块连接对齐位置，文本上下文来自BERT，模型不另设音视跨位置注意力或三方高阶交互。')
head('3 不确定性融合、输出与解释定义')
para('首先从未进行精度缩放的单项计算每模态四维净分数uₘ=β+Σⱼaₘⱼ，同时统计各输出维度局部读出的绝对活性Σⱼ|aₘⱼ|。拼接后得到8维输入，通过每模态一个8→1线性层预测有界对数方差。训练时活性来自模态丢弃后的单项，单模态监督仍使用丢弃前的实际单项。')
eq('vₘ = −1 + 4 tanh(wₘᵀ[uₘ; Σⱼ|aₘⱼ|] + bₘ)； σₘ² = exp(vₘ)')
eq('pₘ = 1ₘ exp(−vₘ) / Σₙ 1ₙ exp(−vₙ)； κₘ = K pₘ，K = Σₘ 1ₘ')
para('1ₘ表示当前保留且可观测的模态。单项按κₘ缩放，成对项按√(κₘκₙ)缩放。精度相同时κₘ=1，恢复基础加性读出；方差被限制在exp(−5)至exp(3)之间。方差监督针对单模态回归残差，它既不是分类概率校准，也不是融合后预测区间。归一化精度pₘ是可靠性系数，不能直接作为模态贡献百分比。')
eq('s = β + Σₘⱼ κₘaₘⱼ + Σₘₙⱼ √(κₘκₙ)iₘₙⱼ； q=s₀； ℓ=(s₁,s₂,s₃)')
para('回归原始输出r=3 tanh(q)。冻结分类策略为中性logit加−0.4后取最大值：ℓ′=ℓ+(0,−0.4,0)，ĉ=argmax ℓ′。最终强度按类别保持符号一致：负向min(r,−10⁻⁶)，中性0，正向max(r,10⁻⁶)。所有主表MAE均评价该最终强度，同时另存raw_intensity。分类和回归共享全部编码及交互参数，输出层有各任务对应的行，联合训练；主模型未启用有序分类头，但保留强度的有序辅助损失。')
para('对每条样本，以预测类别ĉ与最强竞争类别c₂的校准logit差M=ℓ′ĉ−ℓ′c₂为分类解释目标。将所有4维读出投影到这一类别差后，记单模态总项Aₘ、交互总项Iₘₙ。每个成对交互的一半分配给相关两端，得到模态净分数与位置净分数：')
eq('Φₘ = Aₘ + ½ Σₙ≠ₘ Iₘₙ； ψₘⱼ = ãₘⱼ + ½ Σₙ≠ₘ ĩₘₙⱼ； M = βM + ΣₘΦₘ')
eq('ρₘ = |Φₘ| / Σₙ|Φₙ|； principal = argmaxₘ |Φₘ|')
para('ã和ĩ表示已缩放、已投影到解释目标的读出。βM包括两个对应输出偏置之差，以及中性−0.4校准项对该类别差的影响。解释卡同时保留符号：正值支持预测类别相对于竞争类别，负值反对该判断。ρ采用模态净值的绝对值归一化，偏置βM单列；它不同于把所有局部绝对值先相加得到的活性比例。若分母近零则标为未定义，主要模态不能强制指定。回归解释以q为目标同样核算，不能将其比例称为最终非线性强度的线性分配。动态可靠性与BERT上下文均随输入变化，因此当前分数项删除不等于重新运行被修改输入后的输出变化。')
head('4 目标函数、训练方案与关键参数')
eq('L = LHuber(r,y) + 0.5 LCE(ℓ,c) + 0.1 Lord + 0.001 Lpair + 0.1 Luni + 0.1 Lneutral + 0.1 Lpolarity + 0.1 LNLL')
para('Huber阈值δ=1。Lpair为每样本所有已缩放成对项回归分量的绝对值之和，再在批内平均。Luni在实际可观测模态上平均Huber单模态回归损失与0.5倍单模态交叉熵；它使用uₘ的回归、分类分量。Lneutral使用中性logit减去正负logsumexp作为二分类log-odds；Lpolarity仅在非中性样本上对正负logit差进行二元交叉熵。Lord用固定带宽0.3、温度0.2构造两个累计logit：((r+0.3)/0.2, (r−0.3)/0.2)，分别以1[c>0]和1[c>1]为目标计算二元交叉熵并平均；该项不是启用独立的可学习有序头。')
eq('LNLL = mean有效模态{ ½[exp(−vₘ)(3 tanh(uₘ,₀)−y)² + vₘ] }')
para('异方差回归的NLL建模[7]用于单模态残差方差，使可靠性模块获得可监督训练信号。该损失不是三模态贡献的监督真值，预测方差的绝对大小不能直接等同于人类对证据可靠性的判断。整个模型用AdamW更新，BERT编码层和其他层使用固定分组学习率，当前实验没有学习率衰减或warmup。')
table(['参数','冻结主模型设置'],[['随机种子 / 批大小','17 / 32'],['BERT / 其他层学习率','2×10⁻⁵ / 3×10⁻⁴，固定'],['AdamW权重衰减 / 梯度裁剪','0.01（二维以上参数）/ 1.0'],['窗口 / 行投影 / 局部维度 / 交互秩','1 / 64 / 64 / 8'],['局部dropout / 模态丢弃率','0.2 / 0.1'],['最大轮次 / 早停耐心值','20 / 连续4轮三个选优目标均无刷新'],['实际轮数 / 本包检查点','9 / 第5轮'],['可训练BERT层 / AMP','12层，embedding和pooler冻结 / 关闭'],['总参数 / 可训练参数','109,555,007 / 85,127,231'],['分类最终策略','中性logit偏置−0.4；最大logit类别']])
para('每轮仅在验证集搜索策略：441组回归中性阈值和31个中性logit偏置，共472个候选；分别按Macro-F1、Accuracy及最终强度MAE维护独立最佳检查点。任一目标排序刷新即重置早停计数，不是只监测训练损失。本包主表统一取原验证Macro-F1策略，不能混合不同轮次、不同偏置的最好指标。第5轮之后训练损失继续降低但验证原始MAE未刷新，支持保留早停；增加训练上限不自动增加有效训练轮数。')
figure('02_training','图2 主模型训练曲线。虚线为第5轮，实际第9轮停止。右图为原始强度MAE，区别于主表最终强度MAE。')
head('5 开源基线及基础性能评价')
para('选择两个具有真实已训练权重的开源适配基线。SELF-MM[2]采用BERT文本编码、音视LSTM、融合及单模态读出和中心驱动动态伪标签；本实现增加比赛三分类读出、有效观测打包与数值稳定处理。TETFN[3]直接使用MMSA固定版本的网络源码，恢复对齐音视位置，增加比赛分类头及有界回归输出，并使用本地联合训练目标。二者均为比赛数据适配实验，不是原论文数据规模、训练目标和配置下的严格复现。')
para('所有方法仅以3395条训练样本学习，在同一728条验证集上选轮次和策略；本表统一使用各自验证Macro-F1版本。主模型最终按分类logit决策，而两基线的验证最优策略为回归阈值决策，策略差异一并列出，避免把结果误认为统一最后一层阈值的比较。未将NAM、MultiBench仓库或简化TETFN另算作新的已复现独立基线。')
labels={'uncertainty':'不确定性融合（主）','selfmm':'SELF-MM适配','tetfn_source':'TETFN源码适配'}
for split,title in [('valid','验证集（728条）'),('test','附件2测试集（727条）')]:
    table([title,'Acc(%)','Macro-F1','MAE','Pearson','中性F1'],[[labels[x['family']],f"{100*float(x[split+'_accuracy']):.2f}",*[f"{float(x[split+'_'+k]):.4f}" for k in ['macro_f1','mae','pearson','neutral_f1']]] for x in CMP])
table(['模型','检查点轮次','最终策略','附件4 Acc(%)','附件4 MAE'],[[labels[x['family']],x['epoch'],x['policy'],f"{100*float(x['special_accuracy']):.2f}",f"{float(x['special_mae']):.4f}"] for x in CMP])
para('主模型验证Macro-F1并非三者最高：SELF-MM为0.6390，主模型为0.6323。附件2测试中，主模型比SELF-MM准确率高2.48个百分点、Macro-F1高约0.0104，而MAE稍高约0.0040；相比TETFN，测试准确率、Macro-F1和MAE均更优。选择主模型同时考虑其原生可核算解释，不能称其在所有指标上领先。当前为单随机种子的代表配置比较，未给出多种子方差或统计显著性。')
figure('03_confusion','图3 三分类混淆矩阵，行是真值、列是预测。中性与正向之间的混淆较突出。')
table(['验证类别','Precision','Recall','F1','支持数'],[[CN[int(k)],*[f'{v[a]:.4f}' for a in ['precision','recall','f1']],v['support']] for k,v in MET['valid']['final']['per_class'].items()])
figure('04_regression','图4 验证集最终强度散点与残差分布。预测中性置零，形成ŷ=0的水平带。')
head('6 三模态作用差异与局部重要性')
table(['集合/解释目标','文本(%)','语音(%)','视觉(%)'],[[split+'/'+target,*[f"{100*CON[split][key]['mean_shares'][m]:.2f}" for m in 'TAV']] for split in ['valid','test','special'] for key,target in [('classification','分类间隔'),('regression','回归q')]])
para('上述比例先在每条样本内归一化，再对样本等权平均。验证与测试的分类文本占比均约85%，而回归文本占比约68%—71%，语音和视觉对强度的作用相对更高。附件4分类主要参考模态为文本20/20，回归主要模态为文本18/20、视觉2/20（02、14）；02与14的视觉回归净贡献占比分别为48.12%与41.09%。因此“分类主要看文本”不能推导出“强度估计中的视觉无作用”。')
figure('05_contributions','图5 分类贡献、回归贡献与可靠性权重的对照。三者解释不同对象，不能互换。')
para('附件4平均可靠性权重T/A/V为45.07%/28.80%/26.12%，分类净贡献却为89.78%/4.71%/5.51%。这来自不同模态证据大小、方向和抵消程度的共同影响，表明直接把可靠性权重写成贡献会失真。附件4的13号视觉数组全部为零，视觉单项和交互关闭，其视觉贡献应为0；验证和测试分别有15、28条视觉无有效数值观测。')
figure('06_local_16','图6 附件4样本16的主要参考模态内局部重要性。每行给出原BERT位置和对应原文，负值表示反对当前类别间隔；词的上下文由BERT编码。')
figure('06_local_02','图7 附件4样本02的文本局部重要性。保留正负方向，绝对值大不等于支持当前预测。')
head('7 典型样本解释卡与附件4全量结果')
para('解释卡将预测、分类和回归三模态作用、主要模态及关键证据定位放在同一记录中。文本字符区间为原始文本的半开区间；source_index为原BERT序列零基索引。语音时间来自MMS强制对齐，属于机器估计；音视特征没有原始时间戳，视觉只提供同一对齐词段的粗略媒体检索位置，不能断言为生成该特征的精确帧。附件4全20条卡片保存在结果附件中；正文选取02、13、14、16展示模态差异、缺失和误判。')

def evidence(sid,mod):
    e=EXP[sid]; lookup={j:x for x in MEDIA[sid]['entries'] for j in x['source_indices']};lines=[]
    if mod=='V' and SHARE[sid]['available_V']=='False':return '无有效视觉观测；不列视觉证据。'
    for item in e['top_evidence'][mod]:
        if abs(item['signed_score'])<1e-10:continue
        entries=[lookup[j] for j in item['source_indices'] if j in lookup]
        word=' / '.join(dict.fromkeys(z['text'] for z in entries))
        spans=[span for z in entries for span in z.get('char_spans',[])]
        starts=[z.get('start_seconds') for z in entries if z.get('start_seconds') is not None];ends=[z.get('end_seconds') for z in entries if z.get('end_seconds') is not None]
        timing=f'；估计{min(starts):.2f}—{max(ends):.2f}s' if starts and ends else '；时间未定位'
        lines.append(f"位置{item['source_indices']}，{word}，字符{spans}，分数{item['signed_score']:+.4f}"+timing)
    return '；\n'.join(lines) or '无非零局部项。'

notes={'02':'分类与强度使用模态存在差异：分类由文本主导，回归q中视觉净作用最大。局部出现支持与反对证据，不能仅凭单词情感字面解释上下文化读出。','13':'公开标签为正向，模型判为中性。原视觉特征全零，视觉贡献为0；可观测的语言是说明性陈述。只能指出模型缺少视觉数值输入并偏向中性，不能据单个样本证明视觉缺失导致误判。','14':'模型判为中性而公开标签为弱正向。文本主要为人物履历介绍；分类间隔体现中性判断，但回归q中视觉净作用占41.09%。可靠性加权不能消除标签与语义边界的不确定性。','16':'直接负面评价样本。预测与公开标签的极性一致；主要参考文本，但局部值还反映整句上下文，不能把高分词元当作独立因果触发词。'}
card_lines=['# 附件4全20条典型解释卡','本文件只合并已经冻结的模型结果和同一数据的媒体映射；秒数是机器对齐估计，视觉时间不是原特征精确帧。']
for sid in sorted(EXP):
    s=SHARE[sid];r=PRED[sid];lab=LABEL[sid]
    data=[['预测 / 公开源真值',f"{CN[int(r['polarity'])]} {strength(r['intensity'])} / {CN[int(lab['class_id'])]} {strength(lab['sentiment'])}"],['分类贡献T/A/V', ' / '.join(f"{100*float(s['classification_modality_shares_'+m]):.2f}%" for m in 'TAV')],['回归q贡献T/A/V',' / '.join(f"{100*float(s['regression_modality_shares_'+m]):.2f}%" for m in 'TAV')],['分类 / 回归主要模态',s['classification_principal_modality']+' / '+s['regression_principal_modality']],['文本',MEDIA[sid]['raw_text']]]
    card_lines+=['',f'## 样本{sid}']+[f'**{a}：** {b}' for a,b in data]
    for m in 'TAV':card_lines+=[f'**{m}关键证据：** '+evidence(sid,m)]
    if sid in notes:card_lines+=['**解读：** '+notes[sid]]
    if sid in notes:
        add('h2','样本'+sid);table(['项目','内容'],data)
        for m in 'TAV':para(m+'关键证据：'+evidence(sid,m))
        para('解读：'+notes[sid])
(ROOT/'results/附件4_全20条解释卡.md').write_text('\n\n'.join(card_lines)+'\n')
table(['编号','真值','预测','最终强度','T(%)','A(%)','V(%)','主模态'],[[sid,CN[int(LABEL[sid]['class_id'])],CN[int(PRED[sid]['polarity'])],strength(PRED[sid]['intensity']),*[f"{100*float(SHARE[sid]['classification_modality_shares_'+m]):.1f}" for m in 'TAV'],SHARE[sid]['classification_principal_modality']] for sid in sorted(EXP)])
para('附件4错分为01、04、13、14，共4条；其中01、13、14预测为中性，04预测为负向而公开真值为正向。全20条准确率80.00%，最终MAE 0.3390，Pearson 0.9379。样本少、无中性且与附件2测试有重合，这些数值只用于专项案例核验，不能替代727条测试的总体评价。')
head('8 验证集错误归因与解释核算')
para('验证集728条中分类错误254条。真实中性的184条中，89条被判为非中性（23条负向、66条正向），中性召回率51.63%；真实正向338条中67条被判中性，真实负向206条中48条被判中性。中性与正向混淆构成主要误差来源。中性边界校准改善取舍，仍不能解决说明性文本、弱情感与标注差异；继续扩大偏置会同时改变正负样本被归零的比例。')
vc={x['id']:x for x in rows('results/contributions/all_samples_contributions.csv') if x['split']=='valid'}
groups=[]
for name,pred in [('真实负向',lambda x:int(x['true_class'])==0),('真实中性',lambda x:int(x['true_class'])==1),('真实正向',lambda x:int(x['true_class'])==2),('视觉无有效观测',lambda x:vc[x['id']]['available_V']=='False'),('视觉存在有效观测',lambda x:vc[x['id']]['available_V']=='True')]:
    sub=[x for x in valid if pred(x)];error=sum(x['polarity']!=x['true_class'] for x in sub);mae=np.mean([abs(float(x['intensity'])-float(x['true_intensity'])) for x in sub]);groups.append([name,len(sub),error,f'{100*error/len(sub):.2f}%',f'{mae:.4f}'])
table(['验证分组','样本数','分类错误数','错误率','最终MAE'],groups)
para('分组差异同时受标签分布、语句内容和观测质量影响，只能作为错误诊断线索，不构成模态缺失的因果效应。分类贡献在文本上的集中说明模型更依赖语言判别；语音与视觉的较低分类净贡献也可能来自内部正负项抵消，不能视为它们没有可用信息。强度预测的收缩与中性归零会对强情感样本产生较大残差，图4可用于检查这种系统偏差。')
table(['解释检查','验证','测试','附件4'],[['分类间隔最大核算误差',*[f"{CON[s]['classification']['max_accounting_error']:.2e}" for s in ['valid','test','special']]],['回归q最大核算误差',*[f"{CON[s]['regression']['max_accounting_error']:.2e}" for s in ['valid','test','special']]],['逐样本核查数量',728,727,20]])
para('1475条重载前向的raw强度及三个logit与冻结输出逐值一致，分解最大差异处于浮点计算误差范围。该检查说明解释忠实于当前读出的数值，不证明某个证据词是人类情感成因。本主模型没有单独完成随机扰动对照、多种子解释稳定性或人类证据标注评价，故不填入其他历史模型的解释分数，也不声称全局最小充分证据已得到验证。')
head('9 模型局限与可复现实验说明')
para('本文以较小比赛训练集微调BERT，验证集分布不均衡且用于大量候选比较，存在选择偏差；研究结论需要新的数据或独立重复验证。可靠性模块只增加27个参数，复杂度主要来自BERT，并不自动提供贝叶斯后验或校准置信区间。原生加性读出提高了核算透明度，但不能把上下文化局部值解释成独立词效应；音视位置是给定对齐粒度，没有原始时间戳时只能作估计定位。')
para('复现附件包含三个完整冻结检查点、配置、训练统计、决策策略、必要开源源码、环境版本和操作指南。推理可用本地BERT配置初始化后严格加载完整任务权重；从头重训需另行准备公开BERT预训练参数与比赛训练数据。完整BERT研究交付超过题目三题附件合计50MB限制，不能直接声称满足比赛提交容量。本文报告的精度来自完整浮点检查点，尚无经检验满足该容量的压缩模型，不能把不含权重的精简材料当成等价完整复现。')
head('参考文献')
for s in ['[1] Devlin J, Chang M W, Lee K, Toutanova K. BERT: Pre-training of Deep Bidirectional Transformers for Language Understanding. NAACL-HLT, 2019. https://aclanthology.org/N19-1423/', '[2] Yu W, Xu H, Yuan Z, Wu J. Learning Modality-Specific Representations with Self-Supervised Multi-Task Learning for Multimodal Sentiment Analysis. AAAI, 2021. https://arxiv.org/abs/2102.04830', '[3] Wang D, Guo X, Tian Y, Liu J, He L, Luo X. TETFN: A Text Enhanced Transformer Fusion Network for Multimodal Sentiment Analysis. Pattern Recognition, 2023, 136:109259. https://doi.org/10.1016/j.patcog.2022.109259', '[4] Zadeh A B, Liang P P, Poria S, Cambria E, Morency L P. Multimodal Language Analysis in the Wild: CMU-MOSEI Dataset and Interpretable Dynamic Fusion Graph. ACL, 2018. https://aclanthology.org/P18-1208/', '[5] Agarwal R, Melnick L, Frosst N, et al. Neural Additive Models: Interpretable Machine Learning with Neural Nets. NeurIPS, 2021. https://arxiv.org/abs/2004.13912', '[6] Lou Y, Caruana R, Gehrke J, Hooker G. Accurate Intelligible Models with Pairwise Interactions. KDD, 2013. https://doi.org/10.1145/2487575.2487579', '[7] Kendall A, Gal Y. What Uncertainties Do We Need in Bayesian Deep Learning for Computer Vision? NeurIPS, 2017. https://arxiv.org/abs/1703.04977', '[8] thuiar. MMSA: Multimodal Sentiment Analysis Toolkit. https://github.com/thuiar/MMSA. 源码版本a94e65d07fa1ae0d44e552390074b29b0898edfd；SELF-MM和TETFN在本文中均作比赛适配。']:
    para(s)

doc=Document();sec=doc.sections[0];sec.page_height=Cm(29.7);sec.page_width=Cm(21);sec.top_margin=Cm(2);sec.bottom_margin=Cm(2);sec.left_margin=Cm(2.1);sec.right_margin=Cm(2.1)
normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(10.5);normal._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体');normal.paragraph_format.line_spacing=1.2;normal.paragraph_format.space_after=Pt(6)
for name in ['Title','Heading 1','Heading 2']:
    st=doc.styles[name];st.font.name='Calibri';st._element.rPr.rFonts.set(qn('w:eastAsia'),'黑体');st.font.color.rgb=RGBColor.from_string('17384C')
footer=sec.footer.paragraphs[0];footer.alignment=1;field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
md=[]
for kind,value in blocks:
    if kind=='title':doc.add_heading(value,0);md+=['# '+value,'']
    elif kind in ['h','h2']:doc.add_heading(value,1 if kind=='h' else 2);md+=['## '+value if kind=='h' else '### '+value,'']
    elif kind in ['p','eq']:
        p=doc.add_paragraph(value);md+=[value,'']
        if kind=='eq':p.alignment=1;p.paragraph_format.space_after=Pt(9)
    elif kind=='figure':
        name,caption=value;doc.add_picture(str(FIG/(name+'.png')),width=Cm(16.5));p=doc.add_paragraph(caption);p.alignment=1;p.paragraph_format.keep_with_next=False;md += [f'![{caption}](figures/{name}.png)','']
    elif kind=='table':
        headers,data=value;t=doc.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1'
        for cell,text in zip(t.rows[0].cells,headers):cell.text=str(text)
        repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in data:
            for cell,text in zip(t.add_row().cells,row):cell.text=str(text)
        for row in t.rows:
            pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(3)
                    for run in p.runs:run.font.size=Pt(9)
        md+=['| '+' | '.join(map(str,headers))+' |','|'+'|'.join(['---']*len(headers))+'|']+['| '+' | '.join(str(x).replace('\n','<br>').replace('|','/') for x in row)+' |' for row in data]+['']
path=PAPER/'问题三_不确定性融合论文';doc.save(path.with_suffix('.docx'));path.with_suffix('.md').write_text('\n'.join(md)+'\n')
print(json.dumps({'paper':str(path),'figures':len(doc.inline_shapes),'tables':len(doc.tables),'paragraphs':len(doc.paragraphs),'cards':len(EXP)},ensure_ascii=False))
