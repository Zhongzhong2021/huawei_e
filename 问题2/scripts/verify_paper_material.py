"""Record structural checks after a human visual review of every rendered page."""
import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main(root, qa, reviewed):
    reports = root/'reports'
    docx, = reports.glob('*.docx')
    md = docx.with_suffix('.md')
    content, = reports.glob('*_content.json')
    blocks = json.loads(content.read_text(encoding='utf-8'))
    doc = Document(docx)
    pages = len(PdfReader(qa/'round5.pdf').pages)
    assert reviewed == list(range(1, pages+1)), 'Every page must be visually reviewed'
    images = [qa/f'page-{page}.png' for page in reviewed]
    assert all(p.exists() for p in images)
    text = [p.text for p in doc.paragraphs]
    for block in blocks:
        if block['type'] in ['title', 'heading', 'paragraph']:
            assert block['text'] in text
    tables = [b for b in blocks if b['type'] == 'table']
    assert len(tables) == len(doc.tables)
    for block, table in zip(tables, doc.tables):
        assert [[c.text for c in row.cells] for row in table.rows] == [[str(c) for c in row] for row in [block['header']]+block['rows']]
        assert block['caption'] in text
    equations = [b for b in blocks if b['type'] == 'equation']
    assert len(doc.element.findall('.//'+qn('m:oMath'))) == len(equations)
    markdown = md.read_text(encoding='utf-8')
    for eq in equations:
        assert eq['latex'] in markdown
    figures = [b for b in blocks if b['type'] == 'figure']
    assert len(figures) == len(doc.inline_shapes)
    for fig in figures:
        assert fig['caption'] in text and (reports/fig['path']).exists()
    section = doc.sections[0]
    assert abs(section.page_width.cm-21) < .01 and abs(section.page_height.cm-29.7) < .01
    output = root/'audit/document_qa.json'
    assert not output.exists(), 'Do not replace a previous QA record'
    report = {'verified_utc': datetime.now(timezone.utc).isoformat(), 'docx_sha256': digest(docx),
              'markdown_sha256': digest(md), 'structured_content_sha256': digest(content),
              'pages': pages, 'visually_reviewed_pages': reviewed, 'native_equations': len(equations),
              'tables': len(tables), 'figures': len(figures), 'layout': 'A4 Chinese academic black headings and three-line tables',
              'table_cells_match_structured_source': True, 'paragraphs_match_structured_source': True,
              'editable_equations_and_captions_visually_checked': True,
              'render_method': 'Canonical render_docx.py failed: bundled soffice.exe unavailable; hidden read-only Word COM PDF export and bundled Poppler rasterization used',
              'render_pdf_sha256': digest(qa/'round5.pdf'), 'page_png_sha256': {p.name: digest(p) for p in images}}
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(report, ensure_ascii=False))


if __name__ == '__main__':
    p = argparse.ArgumentParser()
    p.add_argument('root', type=Path)
    p.add_argument('qa', type=Path)
    p.add_argument('--reviewed-pages', type=int, nargs='+', required=True)
    a = p.parse_args()
    main(a.root, a.qa, a.reviewed_pages)
