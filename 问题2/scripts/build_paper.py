"""Evidence-bound Chinese paper material, native MathML/OMML equations.

Run DOCX authoring with the bundled Windows Python. Analysis/figures run in WSL.
"""
import csv
import json
import math
import statistics
import sys
from pathlib import Path
from xml.sax.saxutils import escape

STEM='E题问题2局部模态缺失鲁棒预测论文素材'
def read(p):return json.loads(Path(p).read_text(encoding='utf-8'))
def rows(p):
    with Path(p).open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
def fmt(x):return '未定义' if x is None else f'{x:.4f}'
def avg(v):return statistics.mean(v)
def mean_sd(v):return f'{avg(v):.4f} ± {statistics.stdev(v):.4f}' if len(v)>1 else fmt(v[0])
# Minimal transparent MathML constructors. XSLT converts these to native OMML.
def mi(t):return '<mi>'+escape(str(t))+'</mi>'
def mo(t):return '<mo>'+escape(str(t))+'</mo>'
def mn(t):return '<mn>'+str(t)+'</mn>'
def row(*s):return '<mrow>'+''.join(s)+'</mrow>'
def sub(b,s):return '<msub>'+b+s+'</msub>'
def sup(b,s):return '<msup>'+b+s+'</msup>'
def ss(b,s,u):return '<msubsup>'+b+s+u+'</msubsup>'
def frac(a,b):return '<mfrac>'+a+b+'</mfrac>'
def hat(b):return '<mover accent="true">'+b+mo('^')+'</mover>'
def bar(b):return '<mover accent="true">'+b+mo('¯')+'</mover>'
def summ(lo,hi):return '<munderover>'+mo('∑')+lo+hi+'</munderover>'
def par(*x):return row(mo('('),*x,mo(')'))
def v(name,index=None,power=None):
    b=mi(name)
    if index is not None and power is not None:return ss(b,mi(index),mi(power))
    if index is not None:return sub(b,mi(index))
    if power is not None:return sup(b,mi(power))
    return b
def equations():
    return [
      (r'X_i^a\in\mathbb{R}^{T\times74},\quad X_i^v\in\mathbb{R}^{T\times35},\quad T=50',row(v('X','i','a'),mo('∈'),sup(mi('ℝ'),row(mi('T'),mo('×'),mn(74))),mo(','),v('X','i','v'),mo('∈'),sup(mi('ℝ'),row(mi('T'),mo('×'),mn(35))),mo(','),mi('T'),mo('='),mn(50))),
      (r'\widetilde{o}_{ij}^{m}=v_{ij}o_{ij}^{m}(1-d_{ij}^{m})',row(v('õ','ij','m'),mo('='),v('v','ij'),v('o','ij','m'),par(mn(1),mo('−'),v('d','ij','m')))),
      (r'\mu_k^m={\sum_{i,j}o_{ij}^m x_{ijk}^m\over\sum_{i,j}o_{ij}^m}',row(v('μ','k','m'),mo('='),frac(row(sub(mo('∑'),mi('i,j')),v('o','ij','m'),v('x','ijk','m')),row(sub(mo('∑'),mi('i,j')),v('o','ij','m'))))),
      (r'\sigma_k^m=\max\left(\sqrt{{\sum_{i,j}o_{ij}^m(x_{ijk}^m-\mu_k^m)^2\over\sum_{i,j}o_{ij}^m}},10^{-5}\right)',row(v('σ','k','m'),mo('='),mi('max'),par('<msqrt>'+frac(row(sub(mo('∑'),mi('i,j')),v('o','ij','m'),sup(par(v('x','ijk','m'),mo('−'),v('μ','k','m')),mn(2))),row(sub(mo('∑'),mi('i,j')),v('o','ij','m')))+'</msqrt>',mo(','),sup(mn(10),row(mo('−'),mn(5)))))),
      (r'\widetilde{x}_{ijk}^{m}=\widetilde{o}_{ij}^{m}\,\operatorname{clip}((x_{ijk}^{m}-\mu_k^m)/\sigma_k^m,-10,10)',row(v('x̃','ijk','m'),mo('='),v('õ','ij','m'),mi('clip'),par(frac(row(v('x','ijk','m'),mo('−'),v('μ','k','m')),v('σ','k','m')),mo(','),mo('−'),mn(10),mo(','),mn(10)))),
      (r'\operatorname{Attn}(Q,K,V)=\operatorname{softmax}(QK^{\mathsf T}/\sqrt{64}+B)V',row(mi('Attn'),par(mi('Q'),mo(','),mi('K'),mo(','),mi('V')),mo('='),mi('softmax'),par(frac(row(mi('Q'),sup(mi('K'),mi('T'))),'<msqrt>'+mn(64)+'</msqrt>'),mo('+'),mi('B')),mi('V'))),
      (r'\bar h_i^t={\sum_{j=1}^T\widetilde{o}_{ij}^t H_{ij}\over\max(1,\sum_{j=1}^T\widetilde{o}_{ij}^t)}',row(ss(bar(mi('h')),mi('i'),mi('t')),mo('='),frac(row(summ(row(mi('j'),mo('='),mn(1)),mi('T')),v('õ','ij','t'),v('H','ij')),row(mi('max'),par(mn(1),mo(','),summ(row(mi('j'),mo('='),mn(1)),mi('T')),v('õ','ij','t')))))),
      (r'\bar x_i^m={\sum_{j=1}^T\widetilde{o}_{ij}^m\widetilde{x}_{ij}^m\over\max(1,\sum_{j=1}^T\widetilde{o}_{ij}^m)},\quad m\in\{a,v\}',row(ss(bar(mi('x')),mi('i'),mi('m')),mo('='),frac(row(summ(row(mi('j'),mo('='),mn(1)),mi('T')),v('õ','ij','m'),v('x̃','ij','m')),row(mi('max'),par(mn(1),mo(','),summ(row(mi('j'),mo('='),mn(1)),mi('T')),v('õ','ij','m')))),mo(','),mi('m'),mo('∈'),mo('{'),mi('a'),mo(','),mi('v'),mo('}'))),
      (r'a_i^m={\sum_j\widetilde{o}_{ij}^m\over\max(1,\sum_jv_{ij})},\quad z_i=\operatorname{GELU}(W_f[h_i^t;h_i^a;h_i^v;a_i]+b_f)',row(v('a','i','m'),mo('='),frac(row(sub(mo('∑'),mi('j')),v('õ','ij','m')),row(mi('max'),par(mn(1),mo(','),sub(mo('∑'),mi('j')),v('v','ij')))),mo(','),v('z','i'),mo('='),mi('GELU'),par(v('W','f'),mo('['),v('h','i','t'),mo(';'),v('h','i','a'),mo(';'),v('h','i','v'),mo(';'),v('a','i'),mo(']'),mo('+'),v('b','f')))),
      (r'p_i=\operatorname{softmax}(W_cz_i+b_c),\quad\hat c_i=\arg\max_kp_{ik},\quad\hat y_i=3\tanh(w_r^{\mathsf T}z_i+b_r)',row(v('p','i'),mo('='),mi('softmax'),par(v('W','c'),v('z','i'),mo('+'),v('b','c')),mo(','),sub(hat(mi('c')),mi('i')),mo('='),sub(mi('argmax'),mi('k')),v('p','ik'),mo(','),sub(hat(mi('y')),mi('i')),mo('='),mn(3),mi('tanh'),par(ss(mi('w'),mi('r'),mi('T')),v('z','i'),mo('+'),v('b','r')))),
      (r'\mathcal L={1\over N}\sum_{i=1}^N[-\log p_{i,c_i}+H_1(\hat y_i-y_i)]',row(mi('ℒ'),mo('='),frac(mn(1),mi('N')),summ(row(mi('i'),mo('='),mn(1)),mi('N')),mo('['),mo('−'),mi('log'),v('p','i,cᵢ'),mo('+'),v('H','1'),par(sub(hat(mi('y')),mi('i')),mo('−'),v('y','i')),mo(']'))),
      (r'H_1(e)=\begin{cases}e^2/2,&|e|\le1,\\|e|-1/2,&|e|>1.\end{cases}',row(v('H','1'),par(mi('e')),mo('='),mo('{'),'<mtable><mtr><mtd>'+frac(sup(mi('e'),mn(2)),mn(2))+'</mtd><mtd>'+row(mo('|'),mi('e'),mo('|'),mo('≤'),mn(1))+'</mtd></mtr><mtr><mtd>'+row(mo('|'),mi('e'),mo('|'),mo('−'),frac(mn(1),mn(2)))+'</mtd><mtd>'+row(mo('|'),mi('e'),mo('|'),mo('>'),mn(1))+'</mtd></mtr></mtable>')),
      (r'k_i=\min(n_i,\max(1,\lfloor r n_i+0.5\rfloor)),\quad n_i=\sum_jv_{ij}',row(v('k','i'),mo('='),mi('min'),par(v('n','i'),mo(','),mi('max'),par(mn(1),mo(','),mo('⌊'),mi('r'),v('n','i'),mo('+'),mn('0.5'),mo('⌋'))),mo(','),v('n','i'),mo('='),sub(mo('∑'),mi('j')),v('v','ij'))),
      (r'\operatorname{MacroF1}={1\over3}\sum_{c=0}^2{2TP_c\over2TP_c+FP_c+FN_c}',row(mi('MacroF1'),mo('='),frac(mn(1),mn(3)),summ(row(mi('c'),mo('='),mn(0)),mn(2)),frac(row(mn(2),v('TP','c')),row(mn(2),v('TP','c'),mo('+'),v('FP','c'),mo('+'),v('FN','c'))))),
      (r'\operatorname{MAE}={1\over N}\sum_i|\hat y_i-y_i|,\quad S={1\over2}\overline{F1}_{miss}+{1\over2}(1-\overline{MAE}_{miss}/6)',row(mi('MAE'),mo('='),frac(mn(1),mi('N')),sub(mo('∑'),mi('i')),mo('|'),sub(hat(mi('y')),mi('i')),mo('−'),v('y','i'),mo('|'),mo(','),mi('S'),mo('='),frac(mn(1),mn(2)),sub(bar(mi('F1')),mi('miss')),mo('+'),frac(mn(1),mn(2)),par(mn(1),mo('−'),frac(sub(bar(mi('MAE')),mi('miss')),mn(6))))),
      (r'\mathcal L_{student}=\mathcal L+0.5T_d^2D_{KL}(p_T^{(T_d)}\Vert p_S^{(T_d)})+0.25\mathcal L_{reg}+0.1\mathcal L_{feat}',row(v('ℒ','student'),mo('='),mi('ℒ'),mo('+'),mn('0.5'),ss(mi('T'),mi('d'),mn(2)),v('D','KL'),par(v('p','T','(Td)'),mo('∥'),v('p','S','(Td)')),mo('+'),mn('0.25'),v('ℒ','reg'),mo('+'),mn('0.1'),v('ℒ','feat'))),
    ]

def main(root):
    root=Path(root);out=root/'reports';out.mkdir(exist_ok=True)
    summary=read(root/'analysis/summary.json');conf=summary['confirmation'];dev=summary['development'];final=summary['valid']['final'];ref=summary['valid']['reference']
    figures=read(root/'figures/figure_manifest.json');blocks=[];eqs=equations();eqused=[];tabledata={}
    def title(t):blocks.append({'type':'title','text':t})
    def h(t,level=1):
        if level==2 and t.split(' ',1)[0].replace('.','').isdigit():t=t.split(' ',1)[1]
        blocks.append({'type':'heading','level':level,'text':t})
    def p(t):blocks.append({'type':'paragraph','text':t})
    def eq(index):
        latex,xml=eqs[index-1]
        if index==12:xml=xml.replace('<mo>{</mo><mtable>','<mfenced open="{" close=""><mtable>').replace('</mtable></mrow>','</mtable></mfenced></mrow>')
        blocks.append({'type':'equation','number':index,'latex':latex,'mathml':'<math xmlns="http://www.w3.org/1998/Math/MathML" display="block">'+xml+'</math>'});eqused.append(index)
    def table(name,caption,header,body,widths=None):
        blocks.append({'type':'table','caption':caption,'header':header,'rows':body,'widths':widths});tabledata[name]={'caption':caption,'header':header,'rows':body}
    def fig(index):
        r=figures[index-1];number=1+sum(b['type']=='figure' for b in blocks)
        blocks.append({'type':'figure','paper_number':number,'source_id':r['id'],'path':'../figures/'+r['id']+'.png','caption':f'图{number} '+r['caption']})
    def stats(role,scope,key):return [r[role][scope][key] for r in conf['records']]
    stable=conf['stable_improvement'];selected_name=summary['frozen']['model_name']
    title('局部模态缺失条件下的多模态情感预测建模与验证')
    h('摘要')
    p('针对文本、语音与视觉在局部连续区间不可用的情感预测任务，我们建立了编码前缺失屏蔽、预训练文本表示与观测感知汇总相结合的多任务模型。方法将有效位置与模态观测状态分别建模，使用训练集统计量标准化音视频特征，并通过同一网络联合预测负向、中性、正向类别及连续情感强度。训练仅使用赛题附件2，专项推理保持与训练相同的对齐特征接口。')
    if stable:
        p(f"在按视频分组的内部三折开发后，{('完整词表12层模型' if selected_name=='bert12_full_span' else '裁剪词表12层模型')}通过三个随机种子的官方验证确认。完整输入Macro-F1为{mean_sd(stats('metrics','clean','macro_f1'))}，MAE为{mean_sd(stats('metrics','clean','mae'))}；九类连续缺失场景平均Macro-F1为{mean_sd(stats('metrics','selection','mean_missing_macro_f1'))}，MAE为{mean_sd(stats('metrics','selection','mean_missing_mae'))}。相同种子的第二轮FP32基准完整输入Macro-F1为{mean_sd(stats('reference','clean','macro_f1'))}。")
    else:p('新增候选未满足预先固定的稳定性验收条件，因此保留第二轮FP32模型；新增实验作为探索证据，不据此宣称方法提升。')
    p('我们进一步采用固定缺失位置与形状实验、视频分组配对Bootstrap以及类别和回归误差诊断分析模型边界，并完成附件3全部30条样本预测。结果仅支持本赛题划分与预定缺失机制下的结论；附件3没有标签，输出覆盖率不等同于预测准确率。')
    p('关键词：多模态情感预测；局部模态缺失；预训练语言模型；多任务学习；配对分组评估')
    h('1 问题分析与建模假设')
    p('问题2的核心不是在三模态齐全时追求最高分，而是在一个或多个模态的局部连续区间信息不可用时，仍利用其余可见线索输出情感极性和强度，并解释缺失类型、位置及长度变化与性能的关系。局部缺失与整个模态完全不存在不同；本研究的主实验围绕局部缺失展开，全缺失仅用于验证数值与接口安全性。')
    p('在赛题限定的输入条件下，我们采用附件2与附件3的aligned_50版本。对齐意味着不同模态以相同序列位置组织，不能据此认定各位置具有相同持续秒数。附件2未提供逐词时间戳，因此缺失长度以有效特征位置数及比例表述，且不将位置号伪装为原始视频时间。')
    p('建模采用以下可检验假设。第一，字段内相同样本索引与相同对齐位置具有一致的组织含义。第二，依题意，有效区域内音视频整条特征向量为零表示不可观测，单个坐标为零仍是合法观测。第三，在未提供可靠时间戳的条件下，可用固定位置规则研究局部缺失敏感性，但不能宣称完整模拟了真实设备故障。第四，监督标签遵循负值、零值、正值分别对应负向、中性、正向的定义。')
    p('这些假设存在边界：尾部各模态同时为零且缺少分隔词元时，无法可靠区分真实缺失与填充；我们保留审计标记而不猜测原始长度。模型得到的softmax概率未经校准，不能直接解释为真实正确概率。后文同时报告算法构造、实验发现及其限制。')
    h('2 数据组织与数学定义')
    table('data','表1 数据划分与使用规则',['数据部分','样本数','作用'],[['训练集','3395','拟合模型与预处理统计量；内部三折开发'],['验证集','728','固定候选的三种子确认和规律分析'],['测试集','727','模型冻结后的描述性评价'],['附件3','30','无标签专项最终预测']], [.20,.12,.68])
    p('训练数据包含text_bert、audio、vision及标签。文本统一使用text_bert中的词元编号；不在训练中使用768维连续text特征后又于专项推理临时更换输入。训练、验证与测试共4850条样本已完成本地BERT分词编号核验。附件3的text_bert虽为浮点存储，读取时先验证数值有限且为整数，再转为整型编号。')
    eq(1)
    table('symbols','表2 主要符号与维度',['符号','含义','维度或范围'],[['i，j，m','样本、位置、模态索引','m∈{t,a,v}'],['tᵢ','文本词元编号','50；编号0至30521'],['Xᵢᵃ，Xᵢᵛ','语音与视觉输入','50×74，50×35'],['vᵢⱼ','有效位置标记','0或1'],['oᵢⱼᵐ，dᵢⱼᵐ','原观测标记与人工缺失标记','0或1'],['õᵢⱼᵐ','施加缺失后的可用标记','0或1'],['cᵢ，yᵢ','极性标签与强度标签','{0,1,2}；[−3,3]'],['pᵢ，ŷᵢ','三类概率与预测强度','概率单纯形；[−3,3]'],['hᵢᵗ，hᵢᵃ，hᵢᵛ','三模态投影表示','128，64，64'],['aᵢ，zᵢ','三模态可用比例与融合表示','3，128']], [.23,.47,.30])
    p('有效位置通过文本注意力、词元编号及音视频非零观测的并集确定可辨识边界，排除位置0与CLS、SEP特殊词元。区域内部连续全零仍保留其位置，而不压缩序列。观测掩码则分别判断文本注意力与编号有效性、音视频有限性和整向量全零情况。两类掩码分离，使填充、特殊词元与真实缺失具有不同语义。')
    eq(2)
    p('式（2）中的d=1表示人工移除信息，õ=1表示模型实际可使用该模态位置。人工缺失操作在编码之前生效。原观测掩码在标准化前确定并保存，不能因为标准化后某些数值恰好为零而重新判为缺失。')
    h('3 鲁棒多模态预测模型')
    h('3.1 训练集统计量与输入处理',2)
    p('不同模态的数值尺度差异会影响优化。我们仅使用当前拟合子集内可观测位置估计每个音视频坐标的均值与总体标准差；内部三折各自重新估计，官方完整训练后固定一份统计量用于验证、测试与附件3。')
    eq(3);eq(4);eq(5)
    p('式（3）与式（4）的求和仅覆盖拟合子集，m取语音或视觉。标准差下界为10⁻⁵，标准化后固定裁剪至[−10,10]以限制极端数值的影响；该规则在实验前已固定，不依据验证结果改变。不可观测向量置零但另保留掩码，因此模型不会把占位零当成正常测量。')
    h('3.2 编码前屏蔽与文本上下文表示',2)
    p('文本编码器采用本地bert-base-uncased通用预训练权重初始化[1,2]。词元、位置及全零segment嵌入经LayerNorm与dropout送入Transformer；每层为12个注意力头、768维隐状态和3072维前馈层。只保留前50个位置嵌入以匹配赛题输入。词表裁剪版将拟合子集未出现的编号映射为UNK；完整词表版保留预训练模型原有30522个编号，未读取专项样本建立词表。')
    p('对于缺失的普通词元，先将编号置为占位值并从注意力键中屏蔽，再进行上下文编码；CLS与可辨认的SEP保留为结构边界，CLS位置始终开启以避免全屏蔽注意力的数值问题。汇总仅使用实际观测词元，特殊词元与缺失位置不进入均值。位置与边界信息仍然可见，因此本文没有宣称模型对缺失长度本身完全不知情。')
    eq(6)
    p('式（6）给出单个注意力头的核心计算。Q、K、V为线性投影，头宽为64；B在不可用键位置取−10000，其余位置取0。实现使用残差连接、后置LayerNorm和GELU前馈变换。对于缺失位置形成的查询隐状态，不在后续池化中采用，也不允许其作为有效键传递给其他位置。')
    eq(7)
    p('H为最终文本隐状态。式（7）的分母截为至少1避免空序列除零，池化结果经768→128线性层、LayerNorm与GELU；若文本完全不可用，投影后的表示再次乘以可用性指示，使其严格为零。BERT是已有方法，本文的具体实现与缓存原模型在核验样例上的最大绝对差为8.11×10⁻⁶，该检查用于验证实现兼容性，并非预测性能证据。')
    h('3.3 音视频汇总与多模态融合',2)
    eq(8)
    p('语音和视觉采用观测掩码均值汇总，再分别通过74→64与35→64的线性层、LayerNorm和GELU。均值编码在小样本条件下参数量较低、实现稳定，但不显式区分可见片段内部的细粒度顺序，不能将其描述为精细时序建模。完全缺失的模态输出置零。')
    eq(9)
    p('式（9）将128维文本、64维语音、64维视觉和3维可用比例拼接为259维向量，再映射为128维融合表示，训练时使用0.2的dropout。可用比例告知融合层各模态剩余观测量，但此结构是带可用信息的拼接融合，不是显式门控或动态专家模型，也不构成模态重要性的因果解释。')
    eq(10);fig(1)
    p('分类头输出三类softmax概率，回归头通过3tanh限制强度范围。标签约定0、1、2分别为负向、中性、正向。两头共享融合表示但分别监督，因而可能出现类别与强度符号不一致；本研究原样保存两头输出，不以事后阈值修饰预测。')
    h('3.4 联合目标函数与缺失增强',2)
    eq(11);eq(12)
    p('交叉熵利用离散类别监督，Huber损失在小残差区间提供平滑二次惩罚，在大残差区间降低异常误差的梯度增长。两项损失等权，Huber阈值为1。三类未使用额外类别权重，本阶段也未基于官方验证的中性类别表现事后搜索权重或分类阈值。')
    eq(13)
    p('对每个训练样本，以0.7概率选取一个模态施加连续缺失，模态均匀抽样，比例r从0.1、0.3、0.5中均匀抽样。式（13）适用于有效长度nᵢ>0；空序列不额外移除位置。连续性按有效位置的顺序定义，随机起点保证所选片段在边界内。该增强使优化目标包含人工缺失下的监督风险，但不能保证所有缺失条件都改善，其作用由消融实验判断。')
    h('4 求解流程与实验协议')
    table('algorithm_train','算法1 模型训练与选择',['步骤','操作'],[['1','固定官方划分；在训练集内按原视频标识生成三折，各折视频组不重叠。'],['2','各折只以拟合样本估计标准化参数；按候选定义设置裁剪或完整预训练词表。'],['3','初始化预训练编码器和预测头；每个批量在编码前按规则生成局部缺失掩码。'],['4','计算分类与Huber损失，反向传播、梯度裁剪及AdamW更新。'],['5','内部每轮评价完整输入和九类固定缺失场景；按预定分数保存最佳轮次并早停。'],['6','对完成三折且满足保护条件的候选排序；以最佳轮数中位数固定正式训练轮数。'],['7','仅在官方训练集重新拟合，完成三个种子的验证确认；冻结模型后才开展描述性测试。']], [.08,.92])
    table('hyperparameters','表3 固定训练参数',['项目','取值'],[['编码器学习率','2×10⁻⁵'],['投影与预测头学习率','5×10⁻⁴'],['优化器与权重衰减','AdamW；0.01；偏置和归一化参数不衰减'],['有效批量与梯度裁剪','32；梯度范数上限1'],['学习率计划','前10%更新线性预热，其后线性衰减'],['内部训练预算','12层模型最多6轮；连续3轮无改善早停'],['训练与推理精度','CUDA BF16自动混合精度训练；FP32评价；关闭TF32'],['随机种子','三折分组240924；确认42、2026、3407'],['连续缺失增强','概率0.7；三模态与三个比例均匀抽样'],['损失权重','交叉熵∶Huber=1∶1']], [.35,.65])
    p('AdamW采用解耦权重衰减[3]。显存不足时预设微批量从32降为16或8，并以样本数加权累积梯度，保持有效批量32；实际是否触发降批量以日志为准。固定轮数确认期间不逐轮读取官方验证指标选取检查点，训练结束后一次性评价。')
    eq(14);eq(15)
    p('除式（14）与式（15）外，同时报告Accuracy、各类别F1和Pearson相关系数。F1对三类等权求平均，分母为零的类别按0计；Pearson在样本不足或真值、预测恒定时记为未定义并说明原因，不填造数值。内部选择分数S仅是团队预先固定的比较准则，不是比赛官方评分公式。')
    p('验证主场景为文本、语音、视觉分别缺失10%、30%、50%，共9类；每个样本的随机起点由固定种子、样本ID和场景名称的SHA256确定。补充场景采用前部、中部、后部以及分成最多三段的多个短区间，共36类。多个短片段不保证彼此间隔非零，极短样本中可能合并，报告按实际移除位置统计而非假定严格分离。')
    p('所有候选共享场景定义与样本顺序。S取九场景指标的算术平均，不是将场景样本拼接后计算一次F1。完整输入clean表示没有额外人工缺失，仍可能包含原始自然缺失。人工选择的位置若原本不可观测，不会产生额外的信息移除；因此同时记录所选位置数和新增不可观测位置数。')
    p('候选至少在两个内部折提高S，且三折完整输入平均F1下降不超过0.01、MAE增加不超过0.03。比较基准为第二轮两层蒸馏FP32模型，不混用量化模型结果。只确认一个内部最佳候选；三种子至少两组配对通过相同条件，平均结果也通过，才替换基准。官方测试在前期已被观察，本轮冻结后结果只能作描述性评价。内部交叉验证已用于选择，官方验证也已多轮使用，两者均不能表述为全新独立验证。')
    h('5 实验结果与消融分析')
    body=[]
    for e in dev['entries']:
        m=e['summary'];body.append([{'teacher12_compact_reuse':'12层裁剪词表＋缺失训练','bert12_full_span':'12层完整词表＋缺失训练','bert12_full_none':'12层完整词表 无缺失增强'}[e['name']],*[fmt(m[k]) for k in ['clean_f1','clean_mae','missing_f1','missing_mae','score']]])
    m=dev['reference_summary'];body.insert(0,['第二轮两层蒸馏FP32',*[fmt(m[k]) for k in ['clean_f1','clean_mae','missing_f1','missing_mae','score']]])
    table('cv','表4 内部三折平均结果',['模型','完整F1','完整MAE','缺失F1','缺失MAE','S'],body,[.34,.132,.132,.132,.132,.132])
    p('表4中的F1为Macro-F1，各数值为三折指标的算术平均。裁剪词表12层模型复用已完成的相同六轮协议记录；完整词表方案与关闭增强方案均在三折上训练。关闭增强方案预先指定为消融，不因其结果改变主候选集合。三折分别包含1205、1093、1097条验证样本，折间难度差异通过逐折配对展示。')
    byname={e['name']:e for e in dev['entries']}
    if 'bert12_full_span' in byname:
        a=byname['bert12_full_span']['summary'];b=byname['teacher12_compact_reuse']['summary']
        p(f"完整词表相对于裁剪词表12层模型，三折平均S变化为{a['score']-b['score']:+.4f}，完整F1变化为{a['clean_f1']-b['clean_f1']:+.4f}，完整MAE变化为{a['clean_mae']-b['clean_mae']:+.4f}。该比较提供词表保留策略的经验依据；由于词表维度变化也会影响同种子下后续随机初始化序列，且开发阶段只使用一个训练种子，不能将差异完全归因为未见词元本身。")
    if 'bert12_full_none' in byname:
        a=byname['bert12_full_span']['summary'];b=byname['bert12_full_none']['summary'];delta=a['score']-b['score']
        p(f"连续缺失训练相对于不增强方案的三折平均S变化为{delta:+.4f}，缺失平均F1变化为{a['missing_f1']-b['missing_f1']:+.4f}，缺失平均MAE变化为{a['missing_mae']-b['missing_mae']:+.4f}。"+('该结果支持本配置下有限的平均收益，但仍需结合逐折波动，不能宣称显著优于所有设置。' if delta>0 else '该结果未显示连续缺失训练在本配置下具有更高的综合得分，不能把缺失增强描述为已被充分证明有效的创新点。')+' 两组模型均按自己的内部验证分数选最佳轮次，因此这里比较的是包含早停的训练策略，而非同一固定轮数的纯损失扰动。')
    fig(5)
    p('保留的两层实验分别检验文本单模态、多模态输入、连续缺失增强、蒸馏与随机初始化；四层实验检验另一种编码深度。蒸馏模型与四层模型的内部S差距很小，不能仅凭排序宣称其具有显著优势。两层与12层模型训练轮数及训练目标可能不同，跨结构的总差异不应包装成严格的单因素消融。')
    h('5.1 三种子确认与效果区间',2)
    if conf['records']:
        body=[]
        for role,label in [('reference','第二轮FP32基准'),('metrics','本轮候选')]:
            body.append([label,mean_sd(stats(role,'clean','macro_f1')),mean_sd(stats(role,'clean','mae')),mean_sd(stats(role,'selection','mean_missing_macro_f1')),mean_sd(stats(role,'selection','mean_missing_mae'))])
        table('seeds','表5 官方验证集种子均值与样本标准差',['模型','完整F1','完整MAE','缺失F1','缺失MAE'],body,[.24,.19,.19,.19,.19])
        table('paired','表6 三种子配对确认',['种子','ΔS','Δ完整F1','Δ完整MAE','通过'],[[str(r['seed']),f"{r['deltas']['score']:+.4f}",f"{r['deltas']['clean_f1']:+.4f}",f"{r['deltas']['clean_mae']:+.4f}",'是' if r['passed'] else '否'] for r in conf['pairs']],[.12,.22,.22,.25,.19])
    fig(2)
    intervals=[r for r in summary['intervals'] if r['scenario'] in ['clean','missing_average']]
    table('intervals','表7 固定种子模型差异的95% Bootstrap区间',['范围与指标','差值','95%区间'],[[('完整' if r['scenario']=='clean' else '缺失均值')+' '+{'macro_f1':'Macro-F1','mae':'MAE'}[r['metric']],f"{r['difference']:+.4f}",f"[{r['lower']:+.4f}, {r['upper']:+.4f}]"] for r in intervals],[.44,.18,.38])
    boot=read(root/'analysis/bootstrap_protocol.json')
    p(f"区间估计以原视频为重采样单位，在验证集{boot['clusters']}个视频组中有放回抽取相同数量的组，保留组内全部片段及重复次数，共进行2000次。每次对两个模型和所有场景使用完全相同的抽样组，然后重新计算F1与MAE的差值。该配对设计保留组内相关性和模型比较的样本配对。95%区间取差值分布的2.5%与97.5%分位点，不是三个种子均值的置信区间，也未校正此前模型搜索带来的选择偏差。")
    h('5.2 缺失模态与比例的影响',2);fig(3)
    body=[]
    for mod,label in [('text','文本'),('audio','语音'),('vision','视觉')]:
        for rate in [10,30,50]:
            m=final[f'{mod}_{rate}_random'];body.append([label,str(rate)+'%',fmt(m['macro_f1']),f"{m['macro_f1']-final['clean']['macro_f1']:+.4f}",fmt(m['mae']),f"{m['mae']-final['clean']['mae']:+.4f}"])
    table('missing','表8 最终种子42验证集主缺失场景',['模态','比例','F1','ΔF1','MAE','ΔMAE'],body,[.14,.14,.18,.18,.18,.18])
    f1drop={m:final['clean']['macro_f1']-final[f'{m}_50_random']['macro_f1'] for m in ['text','audio','vision']};worst=max(f1drop,key=f1drop.get)
    p(f"在种子42的50%连续缺失条件下，F1下降最大的模态为{ {'text':'文本','audio':'语音','vision':'视觉'}[worst]}，下降{f1drop[worst]:.4f}。这一现象说明当前模型对该模态的依赖更强，但并不说明该模态在所有真实交互任务中天然最重要。若音视频缺失曲线接近完整输入，应同时考虑其均值表示能力有限、文本主导及噪声移除等可能性；本实验不能独立区分这些机制。")
    p('曲线中的局部非单调变化予以保留。移除信息后某一分类指标小幅上升并不自动构成错误，也不等于缺失有益；可能涉及错误边界变化、被移除的噪声与有限样本波动。本文以所有预定场景和两个任务的联合结果判断鲁棒性，不仅选择符合预期的曲线。')
    h('5.3 缺失位置与连续形状的影响',2);fig(4)
    body=[]
    for mod,label in [('text','文本'),('audio','语音'),('vision','视觉')]:
        scenarios=[f'{mod}_{r}_{shape}' for r in [10,30,50] for shape in ['front','middle','back','multi']]
        worsts=min(scenarios,key=lambda q:final[q]['macro_f1']);bests=max(scenarios,key=lambda q:final[q]['macro_f1'])
        body.append([label,worsts,fmt(final[worsts]['macro_f1']),bests,fmt(final[bests]['macro_f1'])])
    table('shape_extrema','表9 预定补充场景中的描述性范围',['模态','最低F1场景','F1','最高F1场景','F1'],body,[.10,.34,.10,.34,.12])
    p('位置影响应在相同模态和比例下比较，而形状影响应比较随机单长片段与multi场景。表9仅用于定位预定网格内的极值，并非经验证的普遍最差或最佳位置；极值是观察后描述，不参与模型选择。正文图5覆盖全部36类场景，逐场景数值和实际缺失率保存在分析数据中。有效长度较短时，至少移除一个位置会使实际比例偏离名义比例；该偏差已单独统计，不能忽略。')
    position_values=[final[f'text_50_{shape}']['macro_f1'] for shape in ['front','middle','back']]
    p(f"固定文本缺失比例为50%时，前部、中部、后部缺失的F1依次为{fmt(position_values[0])}、{fmt(position_values[1])}、{fmt(position_values[2])}。本模型在该验证划分上对后部文本缺失更敏感，但没有原始逐词时间戳及专项定位核验，不能进一步断言后部对应特定情绪转折或结尾词。语音、视觉的位置变化幅度较小，也不能据此推断这些模态完全无用。")
    shape_rows=[]
    for mod,label in [('text','文本'),('audio','语音'),('vision','视觉')]:
        for rate in [10,30,50]:
            a,b=final[f'{mod}_{rate}_random'],final[f'{mod}_{rate}_multi']
            shape_rows.append([label,str(rate)+'%',fmt(a['macro_f1']),fmt(b['macro_f1']),f"{b['macro_f1']-a['macro_f1']:+.4f}",f"{b['mae']-a['mae']:+.4f}"])
    table('span_shape','表10 同等缺失比例下的单长片段与多个短片段',['模态','比例','单段F1','多段F1','ΔF1','ΔMAE'],shape_rows,[.13,.13,.18,.18,.19,.19])
    p('表10的差值均为多个短片段减去随机单长片段。文本的三个比例下，多段缺失的F1均略高，但10%和30%时MAE并未同步下降；音视频的变化方向也不一致。因此，缺失形状的影响依赖模态、比例及评价任务，不能给出“多段一定优于单段”的普遍结论。这些比较使用预定随机起点的一组实现，尚未覆盖缺失起点本身的重复抽样不确定性。')
    actual=[r for r in rows(root/'analysis/actual_missing_rates.csv') if r['scenario'] in ['text_10_random','text_30_random','text_50_random']]
    p('名义缺失率10%、30%、50%对应的样本平均实际位置比例分别为'+ '、'.join(f"{float(r['actual_ratio_mean'])*100:.2f}%" for r in actual)+'。这些差异来自按有效长度取整和至少移除一个位置的规则，而不是修改场景定义；实际新增不可观测位置数另见完整分析表。')
    h('6 分类错误与回归诊断');fig(6)
    cm=final['clean']['confusion_matrix'];neutral_n=sum(cm[1]);d=summary['diagnostics']['final']
    p(f"最终模型完整验证集的负向、中性、正向F1分别为{fmt(final['clean']['class_f1'][0])}、{fmt(final['clean']['class_f1'][1])}、{fmt(final['clean']['class_f1'][2])}。真实中性样本共{neutral_n}条，其中{cm[1][0]}条被判为负向，{cm[1][2]}条被判为正向。中性采用强度严格等于零的定义，不能通过把接近零的标签事后合并为中性来改善指标。")
    p(f"分类头与回归符号的精确一致性检查发现{d['inconsistent_count']}/{d['n']}条输出不一致，占{100*d['inconsistent_rate']:.2f}%。该检查将回归值严格为零才计为中性，因此也包含分类头预测中性但回归头给出小幅非零值的情形；它是接口诊断，不等同于所有这些样本均预测错误。若未来采用一致性损失或阈值联结，需要在独立开发协议下比较，本轮未进行这种事后修改。")
    table('cases','表11 按固定规则选择的验证案例',['样本标识','类型','真类→预测','真实强度','预测强度'],[[c['sample_id'], '错误最大MAE' if c['selection_rule'].startswith('wrong') else '正确中位MAE',f"{c['true_class']}→{c['predicted_class']}",fmt(c['true_intensity']),fmt(c['predicted_intensity'])] for c in summary['cases']],[.34,.22,.14,.15,.15])
    p('案例按每个真实类别分别选择最大MAE的误分类样本与MAE居中的分类正确样本。分类正确不代表强度误差小；反之，轻微跨越零点的预测也可能造成极性错误。此处只根据已有数值证据讨论错误类型，不在缺乏原始视频核查时断言讽刺、语气、否定范围或特定面部表情是原因。')
    h('7 专项预测与结果接口')
    table('algorithm_infer','算法2 附件3统一推理',['步骤','操作'],[['1','按文件名排序读取全部PKL；以文件名和文件内序号构造稳定样本标识。'],['2','执行与训练相同的整数性、形状、有限性和掩码规则检查。'],['3','加载冻结的训练集标准化参数及FP32模型；不在线下载或调用教师。'],['4','按统一接口输出类别、强度和三类概率；核对30条覆盖率与唯一性。'],['5','重载、单条与批量、CPU与GPU结果对照；保存校验记录和摘要。']], [.08,.92])
    prediction=rows(root/'final/attachment3_predictions.csv')
    table('attachment3','表12 附件3全量预测',['文件编号','预测类别','情感强度','最大类概率'],[[r['sample_id'].replace('.pkl::row000',''),{'Negative':'负向','Neutral':'中性','Positive':'正向'}[r['sentiment']],fmt(float(r['intensity'])),fmt(max(float(r['prob_'+c]) for c in ['negative','neutral','positive']))] for r in prediction],[.31,.20,.23,.26])
    p('专项CSV保留sample_id、scenario、class_id、sentiment、intensity、prob_negative、prob_neutral、prob_positive字段。表12为阅读方便保留四位小数，CSV保留计算精度。文件内序号row000不是原始视频ID，最大类概率也不是经过统计校准的置信度。全部30条无标签样本的预测分布未用于调参。')
    h('8 模型复杂度与适用边界')
    p('设序列长T、隐宽d、前馈宽d_ff、层数L。忽略常数与批量因子，文本编码主要计算量为O(L(Td²+T²d+Td·d_ff))；标准注意力实现保存的注意力矩阵规模为O(LhT²)。词表嵌入参数量为Vd，完整词表增加嵌入存储及优化器状态，但不直接增加长度为T的注意力矩阵规模。音视频汇总计算量分别为O(T·74)与O(T·35)，相对于文本编码较小。')
    p('本轮取消文件体积筛选，仅保存FP32研究模型，不据权重压缩推断速度或显存收益。实际耗时受显卡共享占用、批量和数据加载影响，日志中的进程累计峰值不能归因于单一模型。本阶段未开展统一硬件隔离条件下的效率基准，因此不作推理加速结论。')
    p('模型仍有四项重要局限。第一，音视频均值表示忽略细粒度时序，文本缺失较重时可能难以充分补偿。第二，主要场景一次只移除一个模态，尚未系统覆盖多模态同步局部故障。第三，模拟机制是预定位置遮蔽，不代表真实噪声、分布偏移和设备故障的全部形式。第四，样本规模与重复使用验证集限制了泛化证据强度。三种子改善不能保证未知专项样本或比赛排名。')
    h('9 结论')
    p(('本轮在不以附件体积筛选的条件下，确认了较强12层文本编码器作为多模态模型组成部分的稳定收益。' if stable else '本轮未获得符合固定保护条件的稳定升级，因而保留已经核验的第二轮方案。')+' 方法的可复现要点在于编码前严格移除缺失信息、有效位置与观测状态分离、仅拟合训练集预处理统计量，以及使用相同缺失场景配对比较。通过类别与强度的联合评价、缺失位置和形状分析以及原样保留负面消融，结论的适用范围可以明确核验。')
    p('本研究的贡献定位为面向赛题接口的鲁棒预测流程构建与系统实证分析，而不是提出新的预训练基础模型、证明理论最优性或实现因果解释。附件3全量预测已形成统一结果文件；最终准确性仍须以赛方隐藏标签评价为准。')
    h('参考文献')
    refs=[
      '[1] Devlin J, Chang M W, Lee K, Toutanova K. BERT: Pre-training of Deep Bidirectional Transformers for Language Understanding[C]. NAACL-HLT, 2019: 4171–4186. DOI: 10.18653/v1/N19-1423. https://aclanthology.org/N19-1423/',
      '[2] Google. bert-base-uncased model card[EB/OL]. https://huggingface.co/google-bert/bert-base-uncased. 使用本地固定快照86b5e0934494bd15c9632b12f734a8a67f723594，Apache-2.0许可，访问核验日期2026-09-24。',
      '[3] Loshchilov I, Hutter F. Decoupled Weight Decay Regularization[C]. ICLR, 2019. https://arxiv.org/abs/1711.05101.',
      '[4] Hinton G, Vinyals O, Dean J. Distilling the Knowledge in a Neural Network[EB/OL]. arXiv:1503.02531, 2015. https://arxiv.org/abs/1503.02531.',
      '[5] Sun S, Cheng Y, Gan Z, Liu J. Patient Knowledge Distillation for BERT Model Compression[C]. EMNLP-IJCNLP, 2019: 4323–4332. DOI: 10.18653/v1/D19-1441. https://aclanthology.org/D19-1441/']
    for r in refs:p(r)
    h('附录A 蒸馏基准与补充诊断')
    p('第二轮基准抽取预训练BERT第0和第11层构成两层文本编码器，其他汇总与融合结构保持一致。教师为相同拟合子集训练的12层模型；教师和学生接收完全相同的缺失输入，防止教师利用完整文本绕过缺失条件。蒸馏思想参考[4,5]，本实现只使用类别分布、回归输出与最终文本池化表示，不宣称复现逐层Patient Knowledge Distillation。')
    eq(16)
    p('式（16）中温度T_d=2，KL按类别求和再对批量取平均；回归蒸馏为教师与学生强度输出的均方误差，特征蒸馏为各自768维池化表示经过无仿射LayerNorm后的均方误差。教师停止梯度。基准正式训练7轮，新增候选的轮数由内部三折中位数单独确定，因此两者的对比反映整体训练方案差异。')
    fig(7);fig(8)
    h('附录B 冻结后的描述性测试')
    table('test','表B1 已观察测试集上的冻结模型描述性结果',['模型','Accuracy','Macro-F1','MAE','Pearson'],[[label,*[fmt(summary['test_descriptive'][role]['clean'][k]) for k in ['accuracy','macro_f1','mae','pearson']]] for role,label in [('reference','第二轮FP32基准'),('final','冻结最终模型')]],[.28,.18,.18,.18,.18])
    p('测试集727条样本的上述结果来自模型冻结后的统一计算，不用于结构、超参数、阈值、种子或集成权重选择。该测试集在前轮已被查看，不能将本表作为从未触碰留出集的无偏性能估计。验证、测试各46场景与两个模型共184组指标均从逐样本预测独立复算。')
    h('附录C 复现与材料索引')
    p('实验在WSL2 Ubuntu 22.04及RTX 5070上完成。核心环境为Python 3.10.18、PyTorch 2.8.0+CUDA 12.8、NumPy 2.2.6、scikit-learn 1.7.2。原始题目数据未修改，预训练权重来源与SHA256、数据划分、种子、参数、逐轮日志和冻结摘要均保存。推理只需冻结模型、标准化参数与三模态输入，不依赖教师或网络连接。')
    table('artifacts','表C1 实验材料与复现入口',['材料或入口','用途'],[['study/protocol.json 与 frozen.json','固定规则、模型选择及冻结证据'],['runs 与 evidence/round2','本轮日志及保留的旧轮对照记录'],['evaluation 与 analysis','逐样本预测、指标、Bootstrap抽样和案例'],['figures 与 figure_manifest.json','8组矢量与高分辨率图及数据源摘要'],['final/model.pt 与 scaler.npz','最终FP32模型和训练集预处理参数'],['audit/train/evaluate/predict','自动检查、固定配方重训、冻结评价、专项推理'],['analyze/figures/paper/report','统计分析、重绘图形、生成论文材料和完整报告']], [.43,.57])
    p('本机同环境复现与CPU/GPU、单条/批量一致性以audit目录的实际核验记录为准，不承诺跨硬件或跨PyTorch版本逐位一致。论文素材不包含参赛单位、队员姓名或队伍编号。完整研究档案与正式竞赛附件分开管理，正式提交前还需按整队要求处理体积和统一论文格式。')
    md=[]
    for b in blocks:
        if b['type']=='title':md+=['# '+b['text'],'']
        elif b['type']=='heading':md+=['#'*(b['level']+1)+' '+b['text'],'']
        elif b['type']=='paragraph':md+=[b['text'],'']
        elif b['type']=='equation':md+=['$$',b['latex']+rf'\tag{{{b["number"]}}}','$$','']
        elif b['type']=='table':md+=[b['caption'],'','| '+' | '.join(b['header'])+' |','| '+' | '.join(['---']*len(b['header']))+' |']+['| '+' | '.join(map(str,r))+' |' for r in b['rows']]+['']
        elif b['type']=='figure':md+=[f"![{b['caption']}]({b['path']})",'']
    (out/f'{STEM}.md').write_text('\n'.join(md),encoding='utf-8')
    (out/'paper_content.json').write_text(json.dumps(blocks,ensure_ascii=False,indent=2),encoding='utf-8')
    (out/'table_data.json').write_text(json.dumps(tabledata,ensure_ascii=False,indent=2),encoding='utf-8')
    (out/'figure_order.json').write_text(json.dumps([{'paper_number':b['paper_number'],'source_id':b['source_id']} for b in blocks if b['type']=='figure'],ensure_ascii=False,indent=2),encoding='utf-8')
    (out/'equations.tex').write_text('\n\n'.join('\\begin{equation}\n'+eqs[i-1][0]+rf'\tag{{{i}}}'+'\n\\end{equation}' for i in eqused),encoding='utf-8')
    (out/'references_verified.json').write_text(json.dumps({'verified_date':'2026-09-24','references':refs,'primary_sources_only':True},ensure_ascii=False,indent=2),encoding='utf-8')
    try:
        from docx import Document
    except ImportError:
        print('Markdown and structured paper complete; native DOCX requires bundled Windows Python.');return
    build_docx(root,blocks)

def build_docx(root,blocks):
    from docx import Document
    from docx.shared import Cm,Pt,RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from lxml import etree
    stylesheet=Path('C:/Program Files/Microsoft Office/root/Office16/MML2OMML.XSL')
    if not stylesheet.exists():raise FileNotFoundError('Native editable math needs the tested MML2OMML.XSL path')
    transform=etree.XSLT(etree.parse(str(stylesheet)))
    doc=Document();doc.core_properties.author='';doc.core_properties.last_modified_by='';doc.core_properties.title=STEM;doc.core_properties.comments=''
    sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=sec.bottom_margin=Cm(2.0);sec.left_margin=sec.right_margin=Cm(2.0)
    for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
        st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
        fonts=st.element.get_or_add_rPr().get_or_add_rFonts();fonts.set(qn('w:eastAsia'),'宋体' if name in ['Normal','Caption'] else '黑体')
        st.paragraph_format.space_after=Pt(5)
    doc.styles['Normal'].font.size=Pt(11);doc.styles['Normal'].paragraph_format.line_spacing=1.22
    doc.styles['Normal'].paragraph_format.first_line_indent=Pt(22)
    doc.styles['Title'].font.size=Pt(20);doc.styles['Title'].paragraph_format.first_line_indent=Pt(0)
    for name,size in [('Heading 1',14),('Heading 2',12)]:
        doc.styles[name].font.size=Pt(size);doc.styles[name].paragraph_format.space_before=Pt(12);doc.styles[name].paragraph_format.first_line_indent=Pt(0)
    doc.styles['Caption'].font.size=Pt(9);doc.styles['Caption'].font.italic=False;doc.styles['Caption'].font.bold=False;doc.styles['Caption'].paragraph_format.first_line_indent=Pt(0)
    # The bundled Word template can carry a blue Title bottom border.
    for st in doc.styles:
        for border in st.element.xpath('.//w:pBdr'):border.getparent().remove(border)
    def paragraph(text,style=None):
        p=doc.add_paragraph(text,style);p.paragraph_format.widow_control=True;return p
    for b in blocks:
        typ=b['type']
        if typ=='title':paragraph(b['text'],'Title')
        elif typ=='heading':
            p=paragraph(b['text'],'Heading '+str(b['level']))
            if b.get('page_break_before'):p.paragraph_format.page_break_before=True
        elif typ=='paragraph':paragraph(b['text'])
        elif typ=='equation':
            p=paragraph('');p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(8);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            native=transform(etree.fromstring(b['mathml'].encode())).getroot()
            p._p.append(native);r=p.add_run('  ('+str(b['number'])+')');r.font.size=Pt(10)
        elif typ=='figure':
            p=paragraph('');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.keep_with_next=True
            pic=p.add_run().add_picture(str(root/'reports'/b['path']),width=Cm(16.5));pic._inline.docPr.set('descr',b['caption'])
            c=paragraph(b['caption'],'Caption');c.alignment=WD_ALIGN_PARAGRAPH.LEFT;c.paragraph_format.keep_together=True
        elif typ=='table':
            c=paragraph(b['caption'],'Caption');c.alignment=WD_ALIGN_PARAGRAPH.CENTER;c.paragraph_format.keep_with_next=True
            data=[b['header']]+b['rows'];n=len(b['header']);widths=b['widths'] or [1/n]*n
            tab=doc.add_table(rows=0,cols=n);tab.alignment=WD_TABLE_ALIGNMENT.CENTER;tab.autofit=False
            for col,w in zip(tab.columns,widths):col.width=Cm(16.8*w)
            for i,values in enumerate(data):
                r=tab.add_row();r._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                if i==0:r._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
                for j,(cell,value) in enumerate(zip(r.cells,values)):
                    cell.width=Cm(16.8*widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    props=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
                    for side in ['top','bottom','left','right']:
                        edge=OxmlElement('w:'+side);line=(i==0 and side in ['top','bottom']) or (i==len(data)-1 and side=='bottom')
                        edge.set(qn('w:val'),'single' if line else 'nil');edge.set(qn('w:sz'),'8' if i in [0,len(data)-1] and side!='bottom' else '6');edge.set(qn('w:color'),'000000');borders.append(edge)
                    props.append(borders);margins=OxmlElement('w:tcMar')
                    for side in ['top','bottom','left','right']:
                        e=OxmlElement('w:'+side);e.set(qn('w:w'),'65');e.set(qn('w:type'),'dxa');margins.append(e)
                    props.append(margins)
                    p=cell.paragraphs[0];p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_before=p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.08
                    p.paragraph_format.keep_with_next=i==0 or (len(data)<=8 and i<len(data)-1)
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 or len(str(value))>26 else WD_ALIGN_PARAGRAPH.CENTER
                    run=p.add_run(str(value));run.bold=i==0;run.font.size=Pt(9)
            paragraph('').paragraph_format.space_after=Pt(1)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;footer.paragraph_format.first_line_indent=Pt(0)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    for border in doc.element.xpath('.//w:pBdr'):border.getparent().remove(border)
    path=root/'reports'/f'{STEM}.docx';doc.save(path)
    print(path)
if __name__=='__main__':main(Path(sys.argv[1]).resolve())
